"""Exercise a packaged EXE over loopback HTTP, using a complete disposable copy.

Never launch the delivered folder during mutation checks.  The original EXE
folder and the source managed library are hashed before and after.  This is a
delivery smoke test; its answers do not replace the canonical quality report.
"""

from __future__ import annotations

import argparse
from contextlib import closing
import io
import json
import os
from pathlib import Path
import shutil
import socket
import sqlite3
import subprocess
import sys
import time
import traceback
import urllib.error
import urllib.request

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
from scripts.windows_bundle import raw_sha256, write_json  # noqa: E402


def create_fixtures(root: Path) -> dict[str, Path]:
    import fitz
    from docx import Document
    from PIL import Image, ImageDraw, ImageFont
    from pptx import Presentation

    root.mkdir(parents=True, exist_ok=False)
    body = (
        "Release delivery smoke fixture. GSM uses GMSK modulation. "
        "RAKE receivers combine multipath signals to improve reliability. "
        "The marker RELEASESMOKE314 identifies this disposable document."
    )
    paths = {kind: root / ("release_smoke." + suffix) for kind, suffix in (
        ("txt", "txt"), ("markdown", "md"), ("docx", "docx"), ("pptx", "pptx"), ("pdf", "pdf"),
    )}
    paths["txt"].write_text("Release smoke TXT\n\n" + body, encoding="utf-8")
    paths["markdown"].write_text("# Release smoke Markdown\n\n" + body, encoding="utf-8")
    docx = Document()
    docx.add_heading("Release smoke DOCX", level=1)
    docx.add_paragraph(body)
    docx.save(paths["docx"])
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Release smoke PPTX"
    slide.placeholders[1].text = body
    presentation.save(paths["pptx"])
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_textbox(fitz.Rect(50, 50, 540, 500), "Release smoke PDF\n\n" + body, fontsize=14)
    pdf.save(paths["pdf"])
    pdf.close()
    scanned = root / "release_smoke_scanned.pdf"
    image = Image.new("RGB", (1200, 430), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 38)
    draw.multiline_text((35, 35), "Release smoke scanned PDF\nGSM uses GMSK modulation.\nRAKE combines multipath signals.\nRELEASESMOKE314", fill="black", font=font, spacing=28)
    png = io.BytesIO()
    image.save(png, format="PNG")
    pdf = fitz.open()
    page = pdf.new_page(width=1200, height=430)
    page.insert_image(page.rect, stream=png.getvalue())
    pdf.save(scanned)
    pdf.close()
    paths["scanned_pdf_ocr"] = scanned
    return paths


def verify_pristine(bundle: Path) -> dict:
    record = json.loads((bundle / "BUILD_RECORD.json").read_text(encoding="utf-8"))
    actual = {
        path.relative_to(bundle).as_posix(): raw_sha256(path)
        for path in bundle.rglob("*") if path.is_file() and path.name != "BUILD_RECORD.json"
    }
    if actual != record["files"]:
        raise ValueError("Original candidate differs from its completed build record")
    return {"file_count": len(actual), "exe_raw_sha256": record["exe_raw_sha256"], "build_record_raw_sha256": raw_sha256(bundle / "BUILD_RECORD.json")}


class ExeSession:
    def __init__(self, bundle: Path, fixtures: Path):
        self.bundle, self.fixtures = bundle, fixtures
        self.process = None
        self.token = ""
        with socket.socket() as probe:
            probe.bind(("127.0.0.1", 0))
            self.port = probe.getsockname()[1]
        self.base = f"http://127.0.0.1:{self.port}"

    def request(self, method: str, path: str, payload=None, *, authenticated=True, timeout=600):
        data = json.dumps(payload).encode("utf-8") if payload is not None else None
        headers = {"Content-Type": "application/json"}
        if self.token and authenticated:
            headers["Authorization"] = "Bearer " + self.token
        request = urllib.request.Request(self.base + path, data=data, method=method, headers=headers)
        try:
            with urllib.request.urlopen(request, timeout=timeout) as response:
                raw = response.read()
                return response.status, json.loads(raw) if response.headers.get_content_type() == "application/json" else raw.decode("utf-8")
        except urllib.error.HTTPError as error:
            return error.code, json.loads(error.read())

    def start(self):
        environment = dict(os.environ)
        # Explicitly bound test process environment; the bundle's default paths
        # then point to the copied _internal/data, never to the source library.
        for name in ("QA_EMBEDDING_MODEL", "QA_ANSWER_MODEL", "QA_LOG_FILE", "VECTOR_BACKEND", "QA_PAIRING_CODE", "QA_TELEMETRY", "RERANKER_ENABLED"):
            environment.pop(name, None)
        environment["LIBRARY_IMPORT_ROOTS"] = str(self.fixtures)
        self.process = subprocess.Popen(
            [str(self.bundle / "LocalDatabaseQA.exe"), "--host", "127.0.0.1", "--port", str(self.port), "--no-browser"],
            cwd=self.bundle, env=environment, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
            stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
        )
        deadline = time.monotonic() + 120
        while time.monotonic() < deadline:
            if self.process.poll() is not None:
                raise RuntimeError("Packaged process exited during startup: " + str(self.process.returncode))
            try:
                status, bootstrap = self.request("GET", "/api/v2/bootstrap", authenticated=False, timeout=2)
                if status == 200:
                    break
            except (OSError, urllib.error.URLError):
                pass
            time.sleep(0.2)
        else:
            raise TimeoutError("Packaged loopback server did not become ready")
        status, _ = self.request("GET", "/api/v2/health", authenticated=False)
        if status != 401:
            raise AssertionError("Packaged API must require pairing")
        status, _ = self.request("POST", "/api/v2/pair", {"code": "not-a-code"}, authenticated=False)
        if status != 401:
            raise AssertionError("Incorrect pairing code was not rejected")
        status, pair = self.request("POST", "/api/v2/pair", {"code": bootstrap["pairing_code"]}, authenticated=False)
        if status != 200 or not pair.get("token"):
            raise AssertionError("Correct pairing failed")
        self.token = pair["token"]
        return self

    def stop(self):
        if not self.process:
            return
        if self.process.poll() is None and self.token:
            try:
                self.request("POST", "/api/v2/shutdown", {}, timeout=10)
            except OSError:
                pass
        try:
            self.process.wait(timeout=15)
        except subprocess.TimeoutExpired:
            # Only the owned disposable-copy process can reach this fallback.
            self.process.terminate()
            self.process.wait(timeout=10)
            raise RuntimeError("Packaged process failed graceful shutdown")
        if self.process.returncode != 0:
            raise RuntimeError("Packaged process returned nonzero: " + str(self.process.returncode))

    def answer(self, question: str, scope=None):
        payload = {"question": question}
        if scope is not None:
            payload["scope"] = scope
        status, job = self.request("POST", "/api/v2/jobs", payload)
        if status != 202:
            raise AssertionError("Question job was not accepted: " + repr(job))
        status, stream = self.request("GET", "/api/v2/jobs/" + job["job_id"] + "/events")
        events = [json.loads(line) for line in stream.splitlines() if line.strip()]
        final = [event for event in events if event.get("event") == "final"]
        if status != 200 or len(final) != 1:
            raise AssertionError("Question did not finish with one final result: " + repr(events))
        return {"result": final[0]["result"], "event_types": [event.get("event") for event in events]}


def run(bundle: Path, work_root: Path, *, source_library: Path | None = None) -> dict:
    bundle = bundle.resolve(strict=True)
    work_root = work_root.resolve()
    if work_root == bundle or bundle in work_root.parents or work_root in bundle.parents:
        raise ValueError("Smoke work root must be separate from the delivered candidate")
    pristine_before = verify_pristine(bundle)
    source_before = raw_sha256(source_library) if source_library else None
    work_root.mkdir(parents=True, exist_ok=False)
    copied = work_root / "LocalDatabaseQA"
    shutil.copytree(bundle, copied)
    fixtures = create_fixtures(work_root / "fixtures")
    db = copied / "_internal/data/library/documents.sqlite3"
    expected = json.loads((bundle / "build_manifest.json").read_text("utf-8"))
    checks, session = {}, None
    report = {
        "artifact_type": "V4_WINDOWS_EXE_SMOKE_TEST", "checks": checks,
        "candidate_bundle": str(bundle), "tested_disposable_copy": str(copied),
        "exe_raw_sha256": pristine_before["exe_raw_sha256"],
        "production_source_hash": expected["source_identity"]["production_source_hash"],
        "smoke_script_raw_sha256": raw_sha256(Path(__file__)),
        "quality_claim": "FUNCTIONAL_DELIVERY_SMOKE_ONLY_NOT_FULL144",
    }
    try:
        session = ExeSession(copied, work_root / "fixtures")
        session.start()
        status, health = session.request("GET", "/api/v2/health")
        checks["initial_health"] = health
        assert status == 200 and health["documents"] == 6 and health["chunks"] == 637, health
        assert health["answer_model"] == "qwen2.5:7b" and health["embedding_model"] in ("bge-m3", "bge-m3:latest"), health
        assert health["embedding_dimension"] == 1024 and health["vector_backend"] == "sqlite" and health["ok"], health
        checks["pairing_and_authentication"] = "PASS"
        for page in ("/", "/app.js", "/library.html", "/library.js"):
            status, content = session.request("GET", page, authenticated=False)
            assert status == 200 and content, page
        checks["double_click_entry_equivalent_and_static_pages"] = "PASS_EXE_ENTRY_WITH_NO_BROWSER_FLAG"
        nothing = session.answer("GSM使用哪种调制？", {"match_nothing": True})
        assert nothing["result"]["out_of_scope"] and not nothing["result"].get("citations"), nothing
        checks["nothing_scope"] = nothing
        checks["default_question"] = session.answer("GSM 使用 GMSK 调制，请指出来源。")
        assert checks["default_question"]["result"].get("answer"), checks["default_question"]
        status, kb = session.request("POST", "/api/v3/library/knowledge-bases", {"name": "Release smoke disposable KB"})
        assert status == 201, kb
        kb_id = kb["knowledge_base_id"]
        imported = {}
        for kind, path in fixtures.items():
            status, result = session.request("POST", "/api/v3/library/documents", {
                "action": "import", "path": str(path), "knowledge_base_id": kb_id, "tags": ["release-smoke"],
            })
            assert status in (200, 201) and result.get("import_status") == "READY" and result.get("chunk_count", 0) >= 1, (kind, result)
            imported[kind] = result
        checks["actual_five_format_imports_and_scanned_pdf"] = imported
        with closing(sqlite3.connect(db.as_uri() + "?mode=ro", uri=True)) as connection:
            record = connection.execute("SELECT value FROM metadata WHERE key=?", ("import_extras_" + imported["scanned_pdf_ocr"]["document_id"],)).fetchone()
            extras = json.loads(record[0])
            assert extras["import"]["ocr_pages"] >= 1, extras
            checks["packaged_ocr_executed"] = extras
        scoped = session.answer("教材有哪些章节", {"knowledge_base_ids": [kb_id]})
        allowed_ids = {record["document_id"] for record in imported.values()}
        assert scoped["result"].get("citations"), scoped
        assert all(item["document_id"] in allowed_ids for item in scoped["result"].get("citations", [])), scoped
        checks["knowledge_base_scope"] = scoped
        tagged = session.answer("教材有哪些章节", {"tags": ["release-smoke"]})
        assert tagged["result"].get("citations"), tagged
        assert all(item["document_id"] in allowed_ids for item in tagged["result"].get("citations", [])), tagged
        checks["tag_scope"] = tagged
        disabled_id = imported["txt"]["document_id"]
        status, disabled = session.request("POST", "/api/v3/library/documents/" + disabled_id + "/disable", {})
        assert status == 200 and not disabled["enabled"], disabled
        checks["managed_document_disable"] = disabled
        status, library_before_restart = session.request("GET", "/api/v3/library")
        assert status == 200 and library_before_restart["statistics"]["documents"] == 12, library_before_restart
        checks["library_after_mutations"] = library_before_restart["statistics"]
        session.stop()
        session = None
        with closing(sqlite3.connect(db.as_uri() + "?mode=ro", uri=True)) as connection:
            assert connection.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
        before_restart_sha = raw_sha256(db)
        session = ExeSession(copied, work_root / "fixtures")
        session.start()
        status, library_after_restart = session.request("GET", "/api/v3/library")
        assert status == 200 and library_after_restart["statistics"] == library_before_restart["statistics"], library_after_restart
        status, persisted = session.request("GET", "/api/v3/library/documents/" + disabled_id)
        assert status == 200 and not persisted["enabled"], persisted
        checks["restart_preserves_library_and_disabled_state"] = "PASS"
        session.stop()
        session = None
        checks["restart_database_raw_sha256_unchanged"] = before_restart_sha == raw_sha256(db)
        assert checks["restart_database_raw_sha256_unchanged"]
        report["status"] = "PASS"
    except Exception as error:
        report.update({"status": "FAIL", "error": repr(error), "traceback": traceback.format_exc()})
    finally:
        if session is not None:
            try:
                session.stop()
            except Exception as error:
                report.update({"status": "FAIL", "shutdown_error": repr(error)})
        pristine_after = verify_pristine(bundle)
        report["original_candidate_unchanged"] = pristine_before == pristine_after
        report["source_library_unchanged"] = source_before == raw_sha256(source_library) if source_library else "NOT_SUPPLIED"
        if not report["original_candidate_unchanged"] or report["source_library_unchanged"] is False:
            report["status"] = "FAIL"
        write_json(work_root / "smoke_report.json", report)
    return report


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--bundle-root", type=Path)
    parser.add_argument("--work-root", type=Path, required=True)
    parser.add_argument("--source-library", type=Path)
    parser.add_argument("--fixtures-only", action="store_true")
    args = parser.parse_args()
    if args.fixtures_only:
        print(json.dumps({kind: str(path) for kind, path in create_fixtures(args.work_root).items()}, indent=2))
        return 0
    if args.bundle_root is None:
        parser.error("--bundle-root is required unless --fixtures-only is supplied")
    report = run(args.bundle_root, args.work_root, source_library=args.source_library)
    print(json.dumps({"status": report["status"], "report": str(args.work_root / "smoke_report.json"), "error": report.get("error")}, ensure_ascii=False))
    return 0 if report["status"] == "PASS" else 1


if __name__ == "__main__":
    raise SystemExit(main())

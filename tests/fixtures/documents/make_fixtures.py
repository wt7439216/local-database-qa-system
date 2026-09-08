"""Generate the cross-format fixture corpus (Phase C / v3.2).

Produces semantically equivalent small documents in five formats under
tests/fixtures/documents/.  Content is hand-authored public test text with
title, section hierarchy, a repeated term (多径/RAKE), a table, Chinese
unicode and English acronyms (GSM/WCDMA/OFDMA/RAKE) — never user-private
documents.  Requires the ingest dependencies (PyMuPDF, python-docx,
python-pptx); run from the repo root:

    python -X utf8 tests/fixtures/documents/make_fixtures.py
"""

from __future__ import annotations

from pathlib import Path
import sys

FIXTURE_DIR = Path(__file__).resolve().parent

SECTIONS = [
    (
        "1. 多径传播",
        [
            "多径传播是移动信道中信号经多条路径到达接收端的现象，会导致衰落。",
            "GSM 与 WCDMA 系统都必须对抗多径衰落带来的影响。",
            "RAKE 接收机可以合并多径能量，是 CDMA 系列系统的关键部件。",
        ],
    ),
    (
        "2. 抗衰落技术",
        [
            "分集接收通过合并独立衰落的支路来对抗衰落。",
            "均衡技术能够补偿信道失真，OFDMA 是 LTE 采用的多址方式。",
        ],
    ),
    (
        "3. 参数表",
        [
            "下表汇总了本文使用的术语与含义。",
        ],
    ),
]

TABLE = [
    ("术语", "含义"),
    ("多径", "信号经多条路径传播"),
    ("分集", "合并独立衰落支路"),
    ("均衡", "补偿信道失真"),
]


def write_markdown(path: Path) -> None:
    lines = ["# 移动通信测试文档", ""]
    for heading, paragraphs in SECTIONS:
        lines.append(f"## {heading}")
        lines.append("")
        lines.extend(paragraphs)
        lines.append("")
    lines.append("| 术语 | 含义 |")
    lines.append("| --- | --- |")
    lines.extend(f"| {term} | {meaning} |" for term, meaning in TABLE)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_text(path: Path) -> None:
    lines = ["第1章 移动通信测试文档", ""]
    for heading, paragraphs in SECTIONS:
        lines.append(heading)
        lines.append("")
        lines.extend(paragraphs)
        lines.append("")
    lines.append("术语表：")
    lines.extend(f"{term}：{meaning}" for term, meaning in TABLE)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_docx(path: Path) -> None:
    from docx import Document

    document = Document()
    document.core_properties.title = "移动通信测试文档"
    document.add_heading("移动通信测试文档", level=1)
    for heading, paragraphs in SECTIONS:
        document.add_heading(heading, level=2)
        for paragraph in paragraphs:
            document.add_paragraph(paragraph)
    table = document.add_table(rows=len(TABLE), cols=2)
    for row_index, (term, meaning) in enumerate(TABLE):
        table.cell(row_index, 0).text = term
        table.cell(row_index, 1).text = meaning
    document.save(str(path))


def write_pptx(path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "移动通信测试文档"
    body = slide.placeholders[1]
    body.text = "\n".join(
        paragraph for _, paragraphs in SECTIONS[:2] for paragraph in paragraphs
    )
    slide2 = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide2.shapes.title.text = "3. 参数表"
    table_shape = slide2.shapes.add_table(len(TABLE), 2, Inches(0.5), Inches(1.5), Inches(8), Inches(2))
    for row_index, (term, meaning) in enumerate(TABLE):
        table_shape.table.cell(row_index, 0).text = term
        table_shape.table.cell(row_index, 1).text = meaning
    presentation.save(str(path))


def write_pdf(path: Path) -> None:
    import fitz

    cjk = "china-s"  # PyMuPDF built-in simplified-Chinese font
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 90), "移动通信测试文档", fontsize=18, fontname=cjk)
    page.insert_text((72, 130), "1. 多径传播", fontsize=13, fontname=cjk)
    for offset, paragraph in enumerate(SECTIONS[0][1], start=1):
        page.insert_text((72, 130 + offset * 15), paragraph, fontsize=10, fontname=cjk)
    page2 = document.new_page()
    page2.insert_text((72, 90), "2. 抗衰落技术", fontsize=13, fontname=cjk)
    for offset, paragraph in enumerate(SECTIONS[1][1], start=1):
        page2.insert_text((72, 110 + (offset - 1) * 15), paragraph, fontsize=10, fontname=cjk)
    page3 = document.new_page()
    page3.insert_text((72, 90), "3. 参数表", fontsize=13, fontname=cjk)
    for row_index, (term, meaning) in enumerate(TABLE):
        page3.insert_text((72, 110 + row_index * 15), f"{term}：{meaning}", fontsize=10, fontname=cjk)
    document.save(str(path))
    document.close()


def main() -> int:
    if len(sys.argv) > 1 and sys.argv[1] in {"--help", "-h"}:
        print(__doc__)
        return 0
    write_markdown(FIXTURE_DIR / "sample.md")
    write_text(FIXTURE_DIR / "sample.txt")
    write_docx(FIXTURE_DIR / "sample.docx")
    write_pptx(FIXTURE_DIR / "sample.pptx")
    write_pdf(FIXTURE_DIR / "sample.pdf")
    print(f"fixtures written to {FIXTURE_DIR}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
